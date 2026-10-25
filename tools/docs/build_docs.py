#!/usr/bin/env python3
"""Genera la documentación de sIsA en Word (.docx) con los diagramas renderizados.

Para cada docs/NN-*.md:
  - extrae los bloques ```mermaid, los guarda en docs/diagramas/<doc>-<n>.mmd
    y los renderiza a .svg y .png (mermaid-cli, vía npx);
  - arma docs/word/<doc>.docx con títulos, tablas, listas, código y los diagramas como imagen.
También crea docs/word/sIsA-Documentacion-completa.docx con todos los documentos.

Requisitos: Python 3, `pip install python-docx`, Node.js (npx descarga @mermaid-js/mermaid-cli).
Uso (desde la raíz del repo):  python tools/docs/build_docs.py
"""
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
OUT = DOCS / "word"
DIAG = DOCS / "diagramas"
MMDC = "@mermaid-js/mermaid-cli@11.4.2"
TITLE = "sIsA"
ORDER = ["README"] + [p.stem for p in sorted(DOCS.glob("0[1-9]-*.md"))]


def render(mmd: Path) -> Path:
    png = mmd.with_suffix(".png")
    svg = mmd.with_suffix(".svg")
    npx = shutil.which("npx") or shutil.which("npx.cmd") or "npx"
    for target, extra in ((png, ["-s", "2"]), (svg, [])):
        cmd = [npx, "-y", MMDC, "-i", str(mmd), "-o", str(target), "-b", "white", *extra]
        r = subprocess.run(cmd, capture_output=True, text=True, shell=(sys.platform == "win32"))
        if r.returncode != 0:
            raise SystemExit(f"Fallo al renderizar {mmd.name}:\n{r.stdout}\n{r.stderr}")
    return png


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\s][^*]*\*)")


def add_inline(par, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)
        elif part.startswith("[") and "](" in part:
            run = par.add_run(part[1 : part.index("](")])
            run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
            run.underline = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1])
            run.italic = True
        else:
            run = par.add_run(part)
        if size and not (part.startswith("`")):
            run.font.size = Pt(size)


def add_table(doc, rows):
    header = [c.strip() for c in rows[0].strip().strip("|").split("|")]
    body = [r for r in rows[2:]]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_inline(cell.paragraphs[0], h, 9)
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, "1F3A5F")
    for line in body:
        cells = [c.strip() for c in re.split(r"(?<!\\)\|", line.strip().strip("|"))]
        row = t.add_row().cells
        for i in range(len(header)):
            row[i].text = ""
            add_inline(row[i].paragraphs[0], cells[i] if i < len(cells) else "", 9)
    doc.add_paragraph()


def build_doc(doc, stem, md_text, diagram_files):
    """Agrega un documento al objeto docx. Devuelve nada."""
    lines = md_text.splitlines()
    i, n_diag = 0, 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            lang = line[3:].strip()
            j = i + 1
            block = []
            while j < len(lines) and not lines[j].startswith("```"):
                block.append(lines[j])
                j += 1
            if lang == "mermaid":
                png = diagram_files[n_diag]
                n_diag += 1
                doc.add_picture(str(png), width=Inches(6.3))
                cap = doc.add_paragraph(f"Figura: {png.stem}")
                cap.runs[0].italic = True
                cap.runs[0].font.size = Pt(8)
            else:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(block))
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            i = j + 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            j = i
            rows = []
            while j < len(lines) and lines[j].startswith("|"):
                rows.append(lines[j])
                j += 1
            add_table(doc, rows)
            i = j
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
        elif re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\s*\d+\.\s+", "", line))
        elif line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, line.lstrip("> ").strip(), 10)
            for r in p.runs:
                r.italic = True
        elif line.strip():
            add_inline(doc.add_paragraph(), line.strip())
        i += 1


def base_document():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
    return doc


def cover(doc, title, subtitle):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(44)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.font.size = Pt(24)
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    DIAG.mkdir(parents=True, exist_ok=True)
    full = base_document()
    cover(full, "Documentación del proyecto", "Cédula electrónica que sabe más y muestra menos · Platanus Hack 26 Caracas · Track Trust")
    for stem in ORDER:
        md = (DOCS / f"{stem}.md").read_text(encoding="utf-8")
        blocks = re.findall(r"```mermaid\n(.*?)```", md, flags=re.S)
        files = []
        for k, src in enumerate(blocks, 1):
            mmd = DIAG / f"{stem}-{k}.mmd"
            mmd.write_text(src, encoding="utf-8")
            print(f"render {mmd.name}")
            files.append(render(mmd))
        one = base_document()
        build_doc(one, stem, md, files)
        one.save(OUT / f"{stem}.docx")
        build_doc(full, stem, md, files)
        full.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        print(f"ok {stem}.docx")
    full.save(OUT / "sIsA-Documentacion-completa.docx")
    print("ok sIsA-Documentacion-completa.docx")


if __name__ == "__main__":
    main()
